from __future__ import annotations

import argparse
from collections.abc import Iterable, Sequence
import json
from pathlib import Path
import sys
from typing import cast

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "output" / "realtime_detail_xlsx"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Convert realtime detail JSON files into structured XLSX workbooks."
    )
    _ = parser.add_argument(
        "--input",
        required=True,
        help="One detail JSON file or a directory that contains detail JSON files.",
    )
    _ = parser.add_argument(
        "--output-dir",
        default=str(DEFAULT_OUTPUT_DIR),
        help="Directory where generated XLSX files will be written.",
    )
    return parser


def normalize_scalar(value: object) -> object:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def as_dict(value: object) -> dict[str, object]:
    return cast(dict[str, object], value) if isinstance(value, dict) else {}


def as_list(value: object) -> list[object]:
    return cast(list[object], value) if isinstance(value, list) else []


def flatten_rows(rows: Iterable[dict[str, object]]) -> pd.DataFrame:
    normalized_rows = [{key: normalize_scalar(value) for key, value in row.items()} for row in rows]
    return pd.DataFrame(normalized_rows)


def build_overview_rows(payload: dict[str, object]) -> list[dict[str, object]]:
    request = as_dict(payload.get("request", {}))
    summary = as_dict(payload.get("summary", {}))
    return [
        {
            "date_from": request.get("date_from", ""),
            "date_to": request.get("date_to", ""),
            "query_mode": request.get("query_mode", ""),
            "query_scope": request.get("query_scope", ""),
            "from_station_name": request.get("from_station_name", ""),
            "to_station_name": request.get("to_station_name", ""),
            "train_code": request.get("train_code", ""),
            "train_class_name": request.get("train_class_name", ""),
            "matched_train_count": summary.get("matched_train_count", ""),
            "available_train_count": summary.get("available_train_count", ""),
            "cheapest_available_price": summary.get("cheapest_available_price", ""),
            "fastest_duration": summary.get("fastest_duration", ""),
            "successful_days": summary.get("successful_days", ""),
            "failed_days": summary.get("failed_days", ""),
            "total_days": summary.get("total_days", ""),
            "partial_failure_count": len(as_list(payload.get("partial_failures", []))),
            "daily_result_count": len(as_list(payload.get("daily_results", []))),
            "aggregated_result_count": len(as_list(payload.get("aggregated_results", []))),
        }
    ]


def build_daily_summary_rows(payload: dict[str, object]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for daily_index, daily_value in enumerate(as_list(payload.get("daily_results", [])), start=1):
        daily = as_dict(daily_value)
        summary = as_dict(daily.get("summary", {}))
        rows.append(
            {
                "daily_index": daily_index,
                "query_date": daily.get("query_date", ""),
                "matched_train_count": summary.get("matched_train_count", ""),
                "available_train_count": summary.get("available_train_count", ""),
                "cheapest_available_price": summary.get("cheapest_available_price", ""),
                "fastest_duration": summary.get("fastest_duration", ""),
                "failure": normalize_scalar(daily.get("failure")),
            }
        )
    return rows


def build_daily_item_rows(payload: dict[str, object]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for daily_index, daily_value in enumerate(as_list(payload.get("daily_results", [])), start=1):
        daily = as_dict(daily_value)
        for item_index, item_value in enumerate(as_list(daily.get("items", [])), start=1):
            item = as_dict(item_value)
            row = {
                "daily_index": daily_index,
                "item_index": item_index,
                **{key: value for key, value in item.items() if key != "seat_offers"},
                "seat_offer_count": len(as_list(item.get("seat_offers", []))),
            }
            rows.append(row)
    return rows


def build_seat_offer_rows(payload: dict[str, object]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for daily_index, daily_value in enumerate(as_list(payload.get("daily_results", [])), start=1):
        daily = as_dict(daily_value)
        for item_index, item_value in enumerate(as_list(daily.get("items", [])), start=1):
            item = as_dict(item_value)
            base = {
                "daily_index": daily_index,
                "item_index": item_index,
                "query_date": item.get("query_date", ""),
                "train_no": item.get("train_no", ""),
                "train_code": item.get("train_code", ""),
                "from_station_name": item.get("from_station_name", ""),
                "to_station_name": item.get("to_station_name", ""),
                "matched_from_station_name": item.get("matched_from_station_name", ""),
                "matched_to_station_name": item.get("matched_to_station_name", ""),
                "matched_station_name": item.get("matched_station_name", ""),
            }
            for offer_index, offer_value in enumerate(as_list(item.get("seat_offers", [])), start=1):
                offer = as_dict(offer_value)
                rows.append({"offer_index": offer_index, **base, **offer})
    return rows


def iter_leaf_rows(value: object, path: str = "root") -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    if isinstance(value, dict):
        if not value:
            rows.append({"path": path, "value": "{}"})
        for key, child in value.items():
            child_path = f"{path}.{key}"
            rows.extend(iter_leaf_rows(child, child_path))
        return rows
    if isinstance(value, list):
        if not value:
            rows.append({"path": path, "value": "[]"})
        for index, child in enumerate(value):
            child_path = f"{path}[{index}]"
            rows.extend(iter_leaf_rows(child, child_path))
        return rows
    rows.append({"path": path, "value": value})
    return rows


def build_raw_json_rows(payload: dict[str, object]) -> list[dict[str, object]]:
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    return [{"line_number": index, "text": line} for index, line in enumerate(text.splitlines(), start=1)]


def autosize_worksheet(path: Path) -> None:
    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_index, column_cells in enumerate(worksheet.columns, start=1):
            if not column_cells:
                continue
            letter = get_column_letter(column_index)
            max_width = 0
            for cell in column_cells:
                text = "" if cell.value is None else str(cell.value)
                if len(text) > max_width:
                    max_width = len(text)
            worksheet.column_dimensions[letter].width = min(max(max_width + 2, 10), 60)
    workbook.save(path)


def export_detail_json_to_xlsx(input_json: Path, output_dir: Path) -> Path:
    payload = as_dict(json.loads(input_json.read_text(encoding="utf-8")))
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{input_json.stem}.xlsx"

    sheets: list[tuple[str, pd.DataFrame]] = [
        ("overview", flatten_rows(build_overview_rows(payload))),
        ("request", flatten_rows([as_dict(payload.get("request", {}))])),
        ("summary", flatten_rows([as_dict(payload.get("summary", {}))])),
        (
            "partial_failures",
            flatten_rows(as_dict(row) for row in as_list(payload.get("partial_failures", []))),
        ),
        (
            "aggregated_results",
            flatten_rows(as_dict(row) for row in as_list(payload.get("aggregated_results", []))),
        ),
        ("daily_summaries", flatten_rows(build_daily_summary_rows(payload))),
        ("daily_items", flatten_rows(build_daily_item_rows(payload))),
        ("seat_offers", flatten_rows(build_seat_offer_rows(payload))),
        ("json_tree", flatten_rows(iter_leaf_rows(payload))),
        ("raw_json", flatten_rows(build_raw_json_rows(payload))),
    ]

    with pd.ExcelWriter(output_path) as writer:
        for sheet_name, dataframe in sheets:
            frame = dataframe if not dataframe.empty else pd.DataFrame([{}])
            frame.to_excel(writer, index=False, sheet_name=sheet_name)

    autosize_worksheet(output_path)
    return output_path


def resolve_input_files(input_path: Path) -> list[Path]:
    if input_path.is_file():
        return [input_path]
    return sorted(path for path in input_path.glob("*.json") if path.is_file())


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    input_path = Path(args.input).expanduser()
    output_dir = Path(args.output_dir).expanduser()
    files = resolve_input_files(input_path)
    if not files:
        print("no json files found", file=sys.stderr)
        return 1

    for path in files:
        output_path = export_detail_json_to_xlsx(path, output_dir)
        print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
